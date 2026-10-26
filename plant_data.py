"""Read the supplied field workbook without changing its source data.

Coordinates are kept exactly as supplied after numeric conversion. Photos are
representative photographs of a species, never claims about a particular GPS fix.
"""

from __future__ import annotations

import colorsys
from collections import Counter
import hashlib
import json
import math
from pathlib import Path
import re
from statistics import median
import unicodedata
from urllib.parse import quote
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class DataError(ValueError):
    """A workbook cannot be used as a plant observation dataset."""


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
# These only bridge names used in the supplied workbook and photo folders.
# They do not merge observation categories or rewrite botanical identifications.
IMAGE_ALIASES = {
    "kakawate / madre de cacao": "kakawate",
    "guava / bayabas": "guava",
    "katakataka": "kalanchoe pinnata",
}
HEADER_ALIASES = {
    "local name": "name",
    "common name": "name",
    "species": "name",
    "scientific name": "scientific_name",
    "longitude": "longitude",
    "lon": "longitude",
    "lng": "longitude",
    "latitude": "latitude",
    "lat": "latitude",
    "elevation": "elevation",
    "elevation m": "elevation",
    "altitude": "elevation",
}


def _text(value: object) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _key(value: object) -> str:
    return _text(value).casefold()


def _slug(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.casefold()).strip("-") or "plant"


def parse_coordinate(value: object, axis: str) -> float:
    """Parse decimal degrees or DMS and validate the latitude/longitude axis.

    Supports numeric Excel cells, ``8.78224 N``, and ``8° 46' 56.064\" N``.
    Invalid values are rejected instead of guessed or silently swapped.
    """
    if axis not in {"latitude", "longitude"}:
        raise ValueError("axis must be latitude or longitude")
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Missing or invalid {axis}")
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        raw = _text(value).upper().replace("−", "-")
        direction = ""
        if raw and raw[0] in "NSEW":
            direction, raw = raw[0], raw[1:].strip()
        if raw and raw[-1] in "NSEW":
            suffix, raw = raw[-1], raw[:-1].strip()
            if direction and suffix != direction:
                raise ValueError(f"Conflicting {axis} directions")
            direction = suffix
        if direction and direction not in ("NS" if axis == "latitude" else "EW"):
            raise ValueError(f"Direction does not match {axis}")
        # Accept degrees, minutes and seconds, but reject other text and more
        # than three components. Decimal commas are not guessed as coordinates.
        raw = re.sub(r"[°º′'″\"’”:]", " ", raw)
        parts = raw.split()
        if not 1 <= len(parts) <= 3 or any(
            not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", p) for p in parts
        ):
            raise ValueError(f"Invalid {axis}: {_text(value)}")
        numbers = [float(p) for p in parts]
        if any(n < 0 or n >= 60 for n in numbers[1:]):
            raise ValueError(f"Invalid DMS minutes or seconds in {axis}")
        sign = -1 if parts[0].startswith("-") else 1
        result = sign * (abs(numbers[0]) + sum(n / 60 ** i for i, n in enumerate(numbers[1:], 1)))
        if direction:
            if sign < 0 and direction in "NE":
                raise ValueError(f"Conflicting sign and direction in {axis}")
            result = abs(result) * (-1 if direction in "SW" else 1)
    limit = 90 if axis == "latitude" else 180
    if not math.isfinite(result) or abs(result) > limit:
        raise ValueError(f"{axis.capitalize()} must be between {-limit} and {limit}")
    return result


def _color(species_key: str, occupied: set[str]) -> str:
    """Derive repeatable, moderately saturated colors from category names."""
    digest = hashlib.sha256(species_key.encode("utf-8")).digest()
    hue = int.from_bytes(digest[:3], "big") / 16777216
    saturation = 0.52 + digest[3] / 255 * 0.24
    lightness = 0.34 + digest[4] / 255 * 0.13
    for attempt in range(360):
        rgb = colorsys.hls_to_rgb((hue + attempt / 360) % 1, lightness, saturation)
        color = "#" + "".join(f"{round(c * 255):02x}" for c in rgb)
        if color not in occupied:
            occupied.add(color)
            return color
    raise DataError("Unable to assign a distinct species color")


def _photo_folders(assets_path: Path, photo_manifest: Path | None = None) -> dict[str, list[Path]]:
    if photo_manifest is not None:
        try:
            manifest = json.loads(Path(photo_manifest).read_text(encoding="utf-8"))
            if manifest.get("version") != 1 or not isinstance(manifest.get("folders"), dict):
                raise ValueError("Unsupported photo manifest format")
            folders = {}
            for folder, names in manifest["folders"].items():
                if not isinstance(folder, str) or not isinstance(names, list) or not names:
                    raise ValueError("Invalid photo folder")
                photos = []
                for name in names:
                    if not isinstance(name, str):
                        raise ValueError("Invalid photo filename")
                    relative = Path(name)
                    if (relative.is_absolute() or len(relative.parts) != 2 or relative.parts[0] != folder
                            or folder in {".", ".."} or "\\" in name or relative.suffix.lower() != ".webp"):
                        raise ValueError("Invalid photo path")
                    photos.append(assets_path / relative)
                folders[_key(folder)] = photos
            return folders
        except (OSError, ValueError, TypeError, AttributeError) as exc:
            raise DataError(f"Cannot read photo manifest: {exc}. Run python build.py to regenerate it.") from exc
    if not assets_path.is_dir():
        return {}
    return {
        _key(folder.name): sorted(
            (p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS),
            key=lambda p: p.name.casefold(),
        )
        for folder in sorted(assets_path.iterdir(), key=lambda p: p.name.casefold())
        if folder.is_dir()
    }


def _image_url(path: Path, assets_path: Path) -> str:
    return "/media/" + quote(path.relative_to(assets_path).as_posix(), safe="/")


def _distance_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    lat1, lon1, lat2, lon2 = map(math.radians, (lat1, lon1, lat2, lon2))
    a = math.sin((lat2 - lat1) / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2
    return 6371.0088 * 2 * math.asin(min(1, math.sqrt(a)))


def load_dataset(workbook_path: Path, assets_path: Path, photo_manifest: Path | None = None) -> dict:
    """Load all usable worksheets and return the website's JSON-ready dataset.

    Each nonempty source row is one observation, including repeated coordinates.
    Plant categories follow local names in the workbook. Scientific names are
    displayed as supplied, with surrounding/repeated whitespace removed only.
    """
    workbook_path, assets_path = Path(workbook_path), Path(assets_path)
    if not workbook_path.is_file():
        raise DataError(f"Workbook not found: {workbook_path.name}. Place it in the project folder or set DATA_FILE.")
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError, KeyError) as exc:
        raise DataError(f"Cannot read {workbook_path.name}: {exc}") from exc

    observations: list[dict] = []
    categories: dict[str, dict] = {}
    issues: list[dict] = []
    source_sheets: list[str] = []
    skipped_rows = 0
    try:
        for sheet in workbook.worksheets:
            columns: dict[str, int] = {}
            header_found = False
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if not any(v is not None and _text(v) for v in values):
                    continue
                current_columns = {
                    HEADER_ALIASES[_key(value)]: i
                    for i, value in enumerate(values)
                    if _key(value) in HEADER_ALIASES
                }
                if {"name", "latitude", "longitude"}.issubset(current_columns):
                    columns = current_columns
                    if not header_found:
                        source_sheets.append(sheet.title)
                    header_found = True
                    continue  # Also allow repeated headers inside a worksheet.
                if not header_found:
                    if row_number >= 25:
                        break
                    continue

                def cell(field: str) -> object:
                    index = columns.get(field)
                    return values[index] if index is not None and index < len(values) else None

                name = _text(cell("name"))
                try:
                    if not name:
                        raise ValueError("Missing local plant name")
                    latitude = parse_coordinate(cell("latitude"), "latitude")
                    longitude = parse_coordinate(cell("longitude"), "longitude")
                except ValueError as exc:
                    skipped_rows += 1
                    issues.append({"code": "invalid_row", "message": str(exc), "source_sheet": sheet.title, "source_row": row_number})
                    continue

                scientific_name = _text(cell("scientific_name")) or None
                category_key = _key(name)
                if category_key not in categories:
                    categories[category_key] = {"name": name, "scientific_name": scientific_name}
                elif scientific_name and categories[category_key]["scientific_name"] not in {None, scientific_name}:
                    issues.append({"code": "scientific_name_variation", "message": f"{name} has more than one supplied scientific name; each observation retains its source value.", "source_sheet": sheet.title, "source_row": row_number})
                elif scientific_name and categories[category_key]["scientific_name"] is None:
                    categories[category_key]["scientific_name"] = scientific_name

                elevation = None
                if cell("elevation") is not None:
                    try:
                        elevation = float(cell("elevation"))
                        if not math.isfinite(elevation):
                            raise ValueError
                    except (TypeError, ValueError):
                        elevation = None
                        issues.append({"code": "invalid_elevation", "message": "Elevation could not be read; this observation is still mapped.", "source_sheet": sheet.title, "source_row": row_number})
                sheet_id = _slug(sheet.title) + "-" + hashlib.sha256(sheet.title.encode()).hexdigest()[:5]
                observations.append({
                    "id": f"{sheet_id}-{row_number}",
                    "_category_key": category_key,
                    "name": name,
                    "scientific_name": scientific_name,
                    "latitude": latitude,
                    "longitude": longitude,
                    "elevation": elevation,
                    "source_sheet": sheet.title,
                    "source_row": row_number,
                })
            if not header_found:
                issues.append({"code": "unsupported_sheet", "message": f"Skipped worksheet {sheet.title}: no local name, latitude and longitude headers found in the first 25 rows.", "source_sheet": sheet.title})
    finally:
        workbook.close()

    if not source_sheets:
        raise DataError("No mapping worksheet found. Required headers: LOCAL NAME, LONGITUDE, LATITUDE. SCIENTIFIC NAME is optional.")
    if not observations:
        raise DataError("No valid plant locations found. Check that each row has a local name and valid latitude/longitude.")

    folders = _photo_folders(assets_path, photo_manifest)
    species = []
    ids: set[str] = set()
    colors: set[str] = set()
    counts = Counter(o["_category_key"] for o in observations)
    category_photos: dict[str, list[Path]] = {}
    for category_key in sorted(categories):
        category = categories[category_key]
        species_id = _slug(category["name"])
        if species_id in ids:
            species_id += "-" + hashlib.sha256(category_key.encode()).hexdigest()[:6]
        ids.add(species_id)
        folder_key = IMAGE_ALIASES.get(category_key, category_key)
        # Scientific-name folders are useful when the workbook uses a local
        # name, as with katakataka. Never borrow a different species' photos.
        photos = folders.get(folder_key, [])
        category_photos[category_key] = photos
        category["id"] = species_id
        species.append({
            "id": species_id,
            "name": category["name"],
            "scientific_name": category["scientific_name"],
            "color": _color(category_key, colors),
            "count": counts[category_key],
            "image_count": len(photos),
            "thumbnail": (_image_url(photos[0], assets_path).replace("/media/", "/thumbnails/", 1)
                          if photo_manifest else _image_url(photos[0], assets_path)) if photos else None,
            "image_folder": photos[0].parent.name if photos else None,
        })

    photo_indexes: Counter = Counter()
    location_counts = Counter((o["latitude"], o["longitude"]) for o in observations)
    duplicate_counts = Counter((o["_category_key"], o["latitude"], o["longitude"]) for o in observations)
    center_lat = median(o["latitude"] for o in observations)
    center_lon = median(o["longitude"] for o in observations)
    distances = [_distance_km(center_lat, center_lon, o["latitude"], o["longitude"]) for o in observations]
    # This identifies records worth checking, without deleting or relocating
    # valid coordinates. The 5 km floor suits a small field survey, not a
    # geographic boundary or a claim that points outside it are wrong.
    distance_limit = max(5.0, median(distances) * 5)
    for observation, distance in zip(observations, distances):
        category_key = observation.pop("_category_key")
        photos = category_photos[category_key]
        observation["species_id"] = categories[category_key]["id"]
        observation["image"] = _image_url(photos[photo_indexes[category_key] % len(photos)], assets_path) if photos else None
        photo_indexes[category_key] += 1
        observation["location_count"] = location_counts[(observation["latitude"], observation["longitude"])]
        if distance > distance_limit:
            issues.append({"code": "coordinate_outlier", "message": f"{observation['name']} is {distance:.1f} km from the median survey location. Coordinate retained exactly as supplied; verify the source record if needed.", "source_sheet": observation["source_sheet"], "source_row": observation["source_row"]})

    photo_paths = {photo for photos in category_photos.values() for photo in photos}
    unmatched_folders = sorted(name for name, photos in folders.items() if photos and name not in {IMAGE_ALIASES.get(k, k) for k in categories})
    if unmatched_folders:
        issues.append({"code": "unmatched_photo_folders", "message": "Photo folders with no matching source plant category were not used: " + ", ".join(unmatched_folders) + "."})
    return {
        "species": species,
        "observations": observations,
        "bounds": [
            [min(o["latitude"] for o in observations), min(o["longitude"] for o in observations)],
            [max(o["latitude"] for o in observations), max(o["longitude"] for o in observations)],
        ],
        "stats": {
            "observations": len(observations),
            "species": len(species),
            "locations": len(location_counts),
            "photos": len(photo_paths),
            "species_with_photos": sum(bool(s["image_count"]) for s in species),
            "skipped_rows": skipped_rows,
            "overlapping_observations": sum(count - 1 for count in location_counts.values()),
            "duplicate_species_coordinates": sum(count - 1 for count in duplicate_counts.values()),
        },
        "issues": issues,
        "source": {"workbook": workbook_path.name, "sheets": source_sheets},
    }
